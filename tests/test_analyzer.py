"""
ExcelPlorer — Tests for Workbook Analyzer

Tests for backend/core/analyzer.py.
Uses programmatically-generated Excel fixtures — no external test files needed.
"""

import json
import pytest
from pathlib import Path

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

from backend.core.analyzer import WorkbookAnalyzer, AnalysisError
from backend.core.schema import (
    AnalysisResult,
    WorkbookSchema,
    SheetSchema,
    ColumnSchema,
)


# ─── Fixtures ───────────────────────────────────────────────────────────────


@pytest.fixture
def tmp_dir(tmp_path):
    """Provide a temporary directory for test files."""
    return tmp_path


@pytest.fixture
def simple_workbook(tmp_dir) -> Path:
    """
    Create a simple .xlsx test workbook with:
    - 3 columns: Name, Age, Email
    - 5 rows of sample data
    - Data validation on Age (whole number 0-150)
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Products"

    # Headers
    ws["A1"] = "Name"
    ws["B1"] = "Age"
    ws["C1"] = "Email"

    # Data
    data = [
        ("Alice", 30, "alice@example.com"),
        ("Bob", 25, "bob@example.com"),
        ("Charlie", 35, "charlie@test.org"),
        ("Diana", 28, "diana@example.com"),
        ("Eve", 22, "eve@test.org"),
    ]
    for i, (name, age, email) in enumerate(data, start=2):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=age)
        ws.cell(row=i, column=3, value=email)

    # Add data validation on Age column
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="150")
    dv.add("B2:B1000")
    ws.add_data_validation(dv)

    path = tmp_dir / "simple.xlsx"
    wb.save(str(path))
    wb.close()
    return path


@pytest.fixture
def complex_workbook(tmp_dir) -> Path:
    """
    Create a complex .xlsx test workbook with:
    - Data sheet with dropdowns, required markers, comments
    - Instructions sheet (non-data)
    - Hidden reference sheet with dropdown values
    - Merged cells
    - Named ranges
    """
    wb = openpyxl.Workbook()

    # ─── Data Sheet ────────────────────────────────────────────────
    ws_data = wb.active
    ws_data.title = "Listing Data"

    # Headers with required markers
    ws_data["A1"] = "SKU ID*"
    ws_data["B1"] = "Product Title*"
    ws_data["C1"] = "Category"
    ws_data["D1"] = "Price"
    ws_data["E1"] = "Color"
    ws_data["F1"] = "Image URL"
    ws_data["G1"] = "Launch Date"

    # Mark required headers with red fill
    red_fill = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")
    ws_data["A1"].fill = red_fill
    ws_data["B1"].fill = red_fill

    # Add a comment
    ws_data["A1"].comment = Comment("Unique product identifier", "System")

    # Sample data
    ws_data["A2"] = "SKU001"
    ws_data["B2"] = "Blue Widget"
    ws_data["C2"] = "Electronics"
    ws_data["D2"] = 29.99
    ws_data["E2"] = "Blue"
    ws_data["F2"] = "https://example.com/img1.jpg"
    ws_data["G2"] = "2024-01-15"

    ws_data["A3"] = "SKU002"
    ws_data["B3"] = "Red Gadget"
    ws_data["C3"] = "Electronics"
    ws_data["D3"] = 49.99
    ws_data["E3"] = "Red"
    ws_data["F3"] = "https://example.com/img2.jpg"
    ws_data["G3"] = "2024-02-20"

    # ─── Reference Sheet (hidden) ──────────────────────────────────
    ws_ref = wb.create_sheet("References")
    ws_ref.sheet_state = "hidden"

    # Dropdown values for Category
    categories = ["Electronics", "Clothing", "Home Decor", "Books", "Sports"]
    for i, cat in enumerate(categories, start=1):
        ws_ref.cell(row=i, column=1, value=cat)

    # Dropdown values for Color
    colors = ["Red", "Blue", "Green", "Black", "White", "Yellow"]
    for i, color in enumerate(colors, start=1):
        ws_ref.cell(row=i, column=2, value=color)

    # ─── Add data validations ──────────────────────────────────────
    # Category dropdown from reference sheet
    dv_category = DataValidation(
        type="list",
        formula1="=References!$A$1:$A$5",
        allow_blank=True,
    )
    dv_category.add("C2:C1000")
    ws_data.add_data_validation(dv_category)

    # Color — inline dropdown
    dv_color = DataValidation(
        type="list",
        formula1='"Red,Blue,Green,Black,White,Yellow"',
        allow_blank=True,
    )
    dv_color.add("E2:E1000")
    ws_data.add_data_validation(dv_color)

    # Price — decimal validation
    dv_price = DataValidation(
        type="decimal",
        operator="greaterThan",
        formula1="0",
    )
    dv_price.add("D2:D1000")
    ws_data.add_data_validation(dv_price)

    # SKU ID — text length validation (required, no blank)
    dv_sku = DataValidation(
        type="textLength",
        operator="lessThanOrEqual",
        formula1="50",
        allow_blank=False,
    )
    dv_sku.showErrorMessage = True
    dv_sku.errorTitle = "Invalid SKU"
    dv_sku.error = "SKU must be 50 characters or less"
    dv_sku.add("A2:A1000")
    ws_data.add_data_validation(dv_sku)

    # ─── Instructions Sheet ────────────────────────────────────────
    ws_instr = wb.create_sheet("Instructions")
    ws_instr["A1"] = "How to fill this template"
    ws_instr["A3"] = "1. Fill all required fields (marked with *)"
    ws_instr["A4"] = "2. Use dropdown values where available"
    ws_instr["A5"] = "3. Do not modify headers"

    # ─── Merge cells in instructions ───────────────────────────────
    ws_instr.merge_cells("A1:D1")

    # ─── Hide a column in data sheet ───────────────────────────────
    ws_data.column_dimensions["G"].hidden = True

    # ─── Named range ───────────────────────────────────────────────
    from openpyxl.workbook.defined_name import DefinedName
    dn = DefinedName("CategoryList", attr_text="References!$A$1:$A$5")
    wb.defined_names.add(dn)

    # Save
    path = tmp_dir / "complex.xlsx"
    wb.save(str(path))
    wb.close()
    return path


@pytest.fixture
def analyzer():
    """Create a fresh WorkbookAnalyzer instance."""
    return WorkbookAnalyzer()


# ─── Basic Analysis Tests ──────────────────────────────────────────────────


class TestBasicAnalysis:
    """Test basic workbook analysis on a simple workbook."""

    def test_analyze_returns_analysis_result(self, analyzer, simple_workbook):
        result = analyzer.analyze(simple_workbook)
        assert isinstance(result, AnalysisResult)
        assert isinstance(result.schema, WorkbookSchema)

    def test_file_metadata(self, analyzer, simple_workbook):
        result = analyzer.analyze(simple_workbook)
        schema = result.schema
        assert schema.filename == "simple.xlsx"
        assert schema.file_extension == ".xlsx"
        assert schema.file_size > 0
        assert schema.sheet_count == 1

    def test_duration_recorded(self, analyzer, simple_workbook):
        result = analyzer.analyze(simple_workbook)
        assert result.duration_ms > 0

    def test_sheet_detected(self, analyzer, simple_workbook):
        result = analyzer.analyze(simple_workbook)
        assert len(result.schema.sheets) == 1
        sheet = result.schema.sheets[0]
        assert sheet.name == "Products"
        assert sheet.is_visible is True
        assert sheet.is_data_sheet is True

    def test_columns_extracted(self, analyzer, simple_workbook):
        result = analyzer.analyze(simple_workbook)
        sheet = result.schema.sheets[0]
        assert len(sheet.columns) == 3

        col_names = [c.name for c in sheet.columns]
        assert "Name" in col_names
        assert "Age" in col_names
        assert "Email" in col_names

    def test_column_indices(self, analyzer, simple_workbook):
        result = analyzer.analyze(simple_workbook)
        columns = result.schema.sheets[0].columns

        name_col = next(c for c in columns if c.name == "Name")
        assert name_col.index == 1
        assert name_col.letter == "A"

        age_col = next(c for c in columns if c.name == "Age")
        assert age_col.index == 2
        assert age_col.letter == "B"

    def test_sample_values(self, analyzer, simple_workbook):
        result = analyzer.analyze(simple_workbook)
        columns = result.schema.sheets[0].columns

        name_col = next(c for c in columns if c.name == "Name")
        assert len(name_col.sample_values) == 5
        assert "Alice" in name_col.sample_values

    def test_type_inference(self, analyzer, simple_workbook):
        result = analyzer.analyze(simple_workbook)
        columns = result.schema.sheets[0].columns

        name_col = next(c for c in columns if c.name == "Name")
        assert name_col.data_type == "string"

        age_col = next(c for c in columns if c.name == "Age")
        assert age_col.data_type == "integer"  # From validation type="whole"

        email_col = next(c for c in columns if c.name == "Email")
        assert email_col.data_type == "email"

    def test_validation_detected(self, analyzer, simple_workbook):
        result = analyzer.analyze(simple_workbook)
        columns = result.schema.sheets[0].columns

        age_col = next(c for c in columns if c.name == "Age")
        assert len(age_col.validation_rules) > 0
        rule = age_col.validation_rules[0]
        assert rule.type == "whole"
        assert rule.operator == "between"


# ─── Complex Analysis Tests ─────────────────────────────────────────────────


class TestComplexAnalysis:
    """Test workbook analysis with complex features."""

    def test_multi_sheet(self, analyzer, complex_workbook):
        result = analyzer.analyze(complex_workbook)
        assert result.schema.sheet_count == 3

    def test_instructions_not_data_sheet(self, analyzer, complex_workbook):
        result = analyzer.analyze(complex_workbook)
        sheets = {s.name: s for s in result.schema.sheets}
        assert sheets["Instructions"].is_data_sheet is False

    def test_hidden_sheet_detected(self, analyzer, complex_workbook):
        result = analyzer.analyze(complex_workbook)
        sheets = {s.name: s for s in result.schema.sheets}
        assert sheets["References"].is_visible is False

    def test_data_sheet_columns(self, analyzer, complex_workbook):
        result = analyzer.analyze(complex_workbook)
        data_sheet = next(
            s for s in result.schema.sheets if s.name == "Listing Data"
        )
        col_names = [c.name for c in data_sheet.columns]
        assert "SKU ID*" in col_names
        assert "Product Title*" in col_names
        assert "Category" in col_names
        assert "Price" in col_names
        assert "Color" in col_names
        assert "Image URL" in col_names

    def test_required_detection_asterisk(self, analyzer, complex_workbook):
        result = analyzer.analyze(complex_workbook)
        data_sheet = next(
            s for s in result.schema.sheets if s.name == "Listing Data"
        )
        columns = {c.name: c for c in data_sheet.columns}
        assert columns["SKU ID*"].is_required is True
        assert columns["Product Title*"].is_required is True
        assert columns["Category"].is_required is False

    def test_dropdown_inline(self, analyzer, complex_workbook):
        result = analyzer.analyze(complex_workbook)
        data_sheet = next(
            s for s in result.schema.sheets if s.name == "Listing Data"
        )
        color_col = next(c for c in data_sheet.columns if c.name == "Color")
        assert len(color_col.allowed_values) > 0
        assert "Red" in color_col.allowed_values
        assert "Blue" in color_col.allowed_values

    def test_dropdown_formula_ref(self, analyzer, complex_workbook):
        result = analyzer.analyze(complex_workbook)
        data_sheet = next(
            s for s in result.schema.sheets if s.name == "Listing Data"
        )
        cat_col = next(c for c in data_sheet.columns if c.name == "Category")
        assert len(cat_col.allowed_values) > 0
        assert "Electronics" in cat_col.allowed_values
        assert "Clothing" in cat_col.allowed_values

    def test_hidden_column_detected(self, analyzer, complex_workbook):
        result = analyzer.analyze(complex_workbook)
        data_sheet = next(
            s for s in result.schema.sheets if s.name == "Listing Data"
        )
        date_col = next(c for c in data_sheet.columns if c.name == "Launch Date")
        assert date_col.is_hidden is True

    def test_comment_extracted(self, analyzer, complex_workbook):
        result = analyzer.analyze(complex_workbook)
        data_sheet = next(
            s for s in result.schema.sheets if s.name == "Listing Data"
        )
        sku_col = next(c for c in data_sheet.columns if c.name == "SKU ID*")
        assert sku_col.comment is not None
        assert "identifier" in sku_col.comment.lower()

    def test_max_length_detected(self, analyzer, complex_workbook):
        result = analyzer.analyze(complex_workbook)
        data_sheet = next(
            s for s in result.schema.sheets if s.name == "Listing Data"
        )
        sku_col = next(c for c in data_sheet.columns if c.name == "SKU ID*")
        assert sku_col.max_length == 50

    def test_merged_cells_detected(self, analyzer, complex_workbook):
        result = analyzer.analyze(complex_workbook)
        instr_sheet = next(
            s for s in result.schema.sheets if s.name == "Instructions"
        )
        assert len(instr_sheet.merged_ranges) > 0

    def test_named_ranges_extracted(self, analyzer, complex_workbook):
        result = analyzer.analyze(complex_workbook)
        assert len(result.schema.named_ranges) > 0
        names = [nr["name"] for nr in result.schema.named_ranges]
        assert "CategoryList" in names

    def test_type_inference_price(self, analyzer, complex_workbook):
        result = analyzer.analyze(complex_workbook)
        data_sheet = next(
            s for s in result.schema.sheets if s.name == "Listing Data"
        )
        price_col = next(c for c in data_sheet.columns if c.name == "Price")
        assert price_col.data_type == "float"

    def test_type_inference_url(self, analyzer, complex_workbook):
        result = analyzer.analyze(complex_workbook)
        data_sheet = next(
            s for s in result.schema.sheets if s.name == "Listing Data"
        )
        url_col = next(c for c in data_sheet.columns if c.name == "Image URL")
        assert url_col.data_type == "url"


# ─── Serialization Tests ───────────────────────────────────────────────────


class TestSerialization:
    """Test that analysis results serialize to valid JSON."""

    def test_to_dict_produces_json_serializable(self, analyzer, complex_workbook):
        result = analyzer.analyze(complex_workbook)
        result_dict = result.to_dict()

        # Should not raise
        json_str = json.dumps(result_dict, indent=2)
        assert len(json_str) > 0

        # Should roundtrip
        parsed = json.loads(json_str)
        assert parsed["schema"]["filename"] == "complex.xlsx"
        assert len(parsed["schema"]["sheets"]) == 3

    def test_schema_helper_methods(self, analyzer, complex_workbook):
        result = analyzer.analyze(complex_workbook)
        schema = result.schema

        data_sheets = schema.get_data_sheets()
        # Only "Listing Data" is visible + data_sheet
        assert len(data_sheets) >= 1
        assert any(s.name == "Listing Data" for s in data_sheets)

        all_cols = schema.get_all_columns()
        assert len(all_cols) >= 6  # At least the 6 visible columns


# ─── Error Handling Tests ──────────────────────────────────────────────────


class TestErrorHandling:
    """Test error handling in the analyzer."""

    def test_file_not_found(self, analyzer, tmp_dir):
        with pytest.raises(FileNotFoundError):
            analyzer.analyze(tmp_dir / "nonexistent.xlsx")

    def test_invalid_extension(self, analyzer, tmp_dir):
        fake = tmp_dir / "test.csv"
        fake.write_text("a,b,c\n1,2,3")
        with pytest.raises(AnalysisError, match="Unsupported file format"):
            analyzer.analyze(fake)

    def test_corrupted_file(self, analyzer, tmp_dir):
        fake = tmp_dir / "corrupt.xlsx"
        fake.write_bytes(b"this is not an xlsx file")
        with pytest.raises(AnalysisError):
            analyzer.analyze(fake)
